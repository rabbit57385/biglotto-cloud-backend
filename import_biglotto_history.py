import os
import sys
from datetime import datetime, date

import openpyxl
import psycopg

TABLE_NAME = "biglotto_draws"
EXPECTED_HEADERS = ["期別", "開獎日期", "獎號1", "獎號2", "獎號3", "獎號4", "獎號5", "獎號6", "特別號"]


def get_connection():
    database_url = os.getenv("DATABASE_URL")
    if database_url:
        return psycopg.connect(database_url)

    password = os.getenv("BIGLOTTO_DB_PASSWORD")
    if password:
        return psycopg2.connect(
            host="localhost",
            port=5432,
            dbname="biglotto",
            user="postgres",
            password=password,
        )

    raise RuntimeError("找不到 DATABASE_URL 或 BIGLOTTO_DB_PASSWORD")


def normalize_draw_no(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def normalize_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                pass
    raise ValueError(f"無法解析開獎日期: {value!r}")


def normalize_number(value, label):
    if value is None:
        raise ValueError(f"{label} 為空")
    number = int(value)
    if not 1 <= number <= 49:
        raise ValueError(f"{label} 超出 1～49: {number}")
    return number


def load_rows(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    records = []

    for ws in wb.worksheets:
        header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        col = {str(name).strip(): idx for idx, name in enumerate(header) if name is not None}
        missing = [name for name in EXPECTED_HEADERS if name not in col]
        if missing:
            print(f"略過工作表 {ws.title}: 缺少欄位 {', '.join(missing)}")
            continue

        sheet_count = 0
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            draw_no = normalize_draw_no(row[col["期別"]])
            if not draw_no:
                continue

            try:
                draw_date = normalize_date(row[col["開獎日期"]])
                main_numbers = [
                    normalize_number(row[col[f"獎號{i}"]], f"獎號{i}")
                    for i in range(1, 7)
                ]
                special = normalize_number(row[col["特別號"]], "特別號")
            except Exception as exc:
                raise ValueError(f"{ws.title} 第 {row_no} 列資料錯誤: {exc}") from exc

            if len(set(main_numbers)) != 6:
                raise ValueError(f"{ws.title} 第 {row_no} 列主號重複: {main_numbers}")
            if special in main_numbers:
                raise ValueError(f"{ws.title} 第 {row_no} 列特別號與主號重複: {special}")

            records.append((draw_no, draw_date, *main_numbers, special))
            sheet_count += 1

        print(f"讀取 {ws.title}: {sheet_count} 期")

    # 防止 Excel 內同一期重複且內容不同。
    by_draw = {}
    for record in records:
        draw_no = record[0]
        if draw_no in by_draw and by_draw[draw_no] != record:
            raise ValueError(f"Excel 內期別 {draw_no} 有兩筆不同資料")
        by_draw[draw_no] = record

    return sorted(by_draw.values(), key=lambda x: (x[1], x[0]))


def import_rows(records):
    conn = get_connection()
    inserted = 0
    skipped = 0

    sql = f"""
        INSERT INTO {TABLE_NAME}
        (draw_no, draw_date, number1, number2, number3, number4, number5, number6, special_number)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (draw_no) DO NOTHING
    """

    try:
        with conn:
            with conn.cursor() as cur:
                for record in records:
                    cur.execute(sql, record)
                    if cur.rowcount == 1:
                        inserted += 1
                    else:
                        skipped += 1

                cur.execute(f"SELECT COUNT(*) FROM {TABLE_NAME}")
                total = cur.fetchone()[0]
                cur.execute(
                    f"""
                    SELECT draw_no, draw_date, number1, number2, number3,
                           number4, number5, number6, special_number
                    FROM {TABLE_NAME}
                    ORDER BY draw_date DESC, draw_no DESC
                    LIMIT 1
                    """
                )
                latest = cur.fetchone()
    finally:
        conn.close()

    return inserted, skipped, total, latest


def main():
    xlsx_path = sys.argv[1] if len(sys.argv) >= 2 else "biglotto_history.xlsx"
    if not os.path.exists(xlsx_path):
        raise FileNotFoundError(f"找不到 Excel: {xlsx_path}")

    print("=== 大樂透歷史資料匯入 ===")
    print(f"Excel: {xlsx_path}")
    records = load_rows(xlsx_path)
    print(f"Excel 有效期數: {len(records)}")

    if not records:
        raise RuntimeError("Excel 沒有可匯入的大樂透資料")

    inserted, skipped, total, latest = import_rows(records)
    print(f"新增: {inserted} 期")
    print(f"已存在/跳過: {skipped} 期")
    print(f"資料庫目前總期數: {total}")
    if latest:
        nums = " ".join(f"{n:02d}" for n in latest[2:8])
        print(f"資料庫最新: {latest[0]} / {latest[1]} / {nums} / 特別號 {latest[8]:02d}")
    print("=== 匯入完成 ===")


if __name__ == "__main__":
    main()
